#!/usr/bin/env python3
"""Build a self-contained HTML workload dashboard from a commitment workbook.

The dashboard is designed for locally shared planning workbooks.  It embeds no
external libraries or data services, so opening the output HTML file does not
upload workbook data anywhere.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
from collections import Counter, defaultdict
from pathlib import Path

from analyze_commitment import analyze


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def dashboard_data(report_dir: Path, threshold: float) -> dict:
    person_months = read_csv(report_dir / "person_month_totals.csv")
    contributions = read_csv(report_dir / "allocation_contributions.csv")
    missing_emails = read_csv(report_dir / "rows_without_email.csv")

    total_by_person = defaultdict(float)
    months = set()
    people = {}
    projects = set()
    for row in person_months:
        row["total_allocation"] = float(row["total_allocation"])
        row["is_overloaded"] = row["total_allocation"] > threshold + 1e-9
        months.add(row["month"])
        people[row["email"]] = row["names_in_workbook"] or row["email"]
        if row["is_overloaded"]:
            total_by_person[row["email"]] += row["total_allocation"] - threshold
    for row in contributions:
        row["allocation"] = float(row["allocation"])
        projects.add(row["project"])

    return {
        "threshold": threshold,
        "personMonths": person_months,
        "contributions": contributions,
        "missingEmails": missing_emails,
        "months": sorted(months),
        "people": [
            {"email": email, "name": people[email], "excess": round(total_by_person[email], 6)}
            for email in sorted(people, key=lambda item: (people[item].lower(), item))
        ],
        "projects": sorted(projects),
    }


def safe_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False).replace("<", "\\u003c")


def build_html(data: dict, workbook_name: str) -> str:
    title = html.escape(f"Commitment workload dashboard — {workbook_name}")
    embedded_data = safe_json(data)
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    :root {{ color-scheme: light; --ink:#172033; --muted:#64748b; --line:#dbe3ee; --blue:#1d4ed8; --red:#b42318; --amber:#b45309; --pale:#f8fafc; }}
    * {{ box-sizing:border-box; }} body {{ margin:0; background:var(--pale); color:var(--ink); font:14px/1.45 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; }}
    header {{ padding:28px max(22px,calc((100vw - 1280px)/2)); background:#0f172a; color:#f8fafc; }}
    h1 {{ margin:0; font-size:24px; }} header p {{ margin:6px 0 0; color:#cbd5e1; }} main {{ max-width:1280px; margin:0 auto; padding:22px; }}
    .notice {{ background:#fffbeb; color:#78350f; border:1px solid #fde68a; padding:11px 13px; border-radius:8px; margin-bottom:18px; }}
    .filters,.tabs {{ display:flex; flex-wrap:wrap; gap:10px; align-items:end; margin-bottom:18px; }} label {{ display:grid; gap:4px; color:#475569; font-weight:650; }} select,input {{ min-width:180px; border:1px solid #cbd5e1; border-radius:6px; padding:8px; background:white; color:var(--ink); }}
    button {{ cursor:pointer; border:1px solid #cbd5e1; border-radius:7px; padding:8px 12px; background:white; color:#334155; font-weight:650; }} button.active {{ background:var(--blue); border-color:var(--blue); color:white; }}
    .cards {{ display:grid; grid-template-columns:repeat(4,minmax(150px,1fr)); gap:12px; margin-bottom:18px; }} .card {{ background:white; border:1px solid var(--line); border-radius:9px; padding:15px; }} .card .label {{ color:var(--muted); font-size:12px; }} .card .value {{ font-size:25px; font-weight:750; margin-top:4px; }}
    .panel {{ background:white; border:1px solid var(--line); border-radius:9px; padding:16px; overflow:auto; }} .hidden {{ display:none; }} h2 {{ font-size:17px; margin:0 0 12px; }}
    table {{ width:100%; border-collapse:collapse; min-width:720px; }} th,td {{ text-align:left; padding:9px 10px; border-bottom:1px solid var(--line); vertical-align:top; }} th {{ color:#475569; font-size:12px; text-transform:uppercase; letter-spacing:.04em; }} tr.over {{ background:#fff1f2; }} .total {{ font-variant-numeric:tabular-nums; font-weight:700; }} .danger {{ color:var(--red); }} .muted {{ color:var(--muted); }} .bar {{ height:9px; min-width:4px; background:#2563eb; border-radius:4px; }} .bar.over {{ background:#dc2626; }} .bar-cell {{ width:230px; }}
    .empty {{ color:var(--muted); padding:22px 0; }} .tag {{ display:inline-block; background:#eff6ff; color:#1e40af; padding:2px 6px; border-radius:4px; margin:1px; font-size:12px; }}
    .matrix {{ min-width:max-content; }} .matrix th {{ position:sticky; top:0; background:white; z-index:2; }} .matrix th:first-child,.matrix td:first-child {{ position:sticky; left:0; background:inherit; z-index:1; min-width:230px; }} .matrix th:first-child {{ z-index:3; }} .matrix td {{ min-width:82px; text-align:center; font-variant-numeric:tabular-nums; }} .matrix td.person {{ text-align:left; background:white; }} .matrix td.load {{ background:#eff6ff; font-weight:700; cursor:help; }} .matrix td.load.watch {{ background:#fef3c7; color:#92400e; }} .matrix td.load.over {{ background:#fee2e2; color:#b42318; }} .matrix td.no-load {{ color:#94a3b8; }}
    .matrix-tooltip {{ position:fixed; display:none; z-index:10; max-width:410px; padding:10px 12px; background:#0f172a; color:#f8fafc; border-radius:7px; box-shadow:0 10px 25px rgba(15,23,42,.22); font-size:12px; line-height:1.45; pointer-events:none; }} .matrix-tooltip .heading {{ font-weight:750; margin-bottom:6px; }} .matrix-tooltip .assignment {{ color:#dbeafe; }}
    @media(max-width:760px) {{ .cards {{ grid-template-columns:repeat(2,minmax(130px,1fr)); }} main {{ padding:14px; }} header {{ padding:20px 14px; }} }}
  </style>
</head>
<body>
<header><h1>{title}</h1><p>Refresh this file by running the generator again with the latest Excel workbook.</p></header>
<main>
  <div class="notice">Workload is aggregated by engineer email. Red rows are strictly above <span id="threshold"></span>. TBH/TBD rows without a parseable email are listed separately.</div>
  <section class="filters">
    <label>Escalation Path<select id="escalationFilter"><option value="">All escalation paths</option></select></label>
    <label>Person<select id="personFilter"><option value="">All people</option></select></label>
    <label>Month<select id="monthFilter"><option value="">All months</option></select></label>
    <label>Project<select id="projectFilter"><option value="">All projects</option></select></label>
    <label><input id="overOnly" type="checkbox"> Overloaded only</label>
    <button id="reset">Reset filters</button>
  </section>
  <section class="cards" id="cards"></section>
  <section class="tabs"><button data-tab="people" class="active">People workload</button><button data-tab="matrix">Workload matrix</button><button data-tab="projects">Project assignments</button><button data-tab="missing">TBH / missing email</button></section>
  <section id="people" class="panel tab"></section>
  <section id="matrix" class="panel tab hidden"></section>
  <section id="projects" class="panel tab hidden"></section>
  <section id="missing" class="panel tab hidden"></section>
</main>
<div id="matrixTooltip" class="matrix-tooltip" role="tooltip"></div>
<script>
const DATA = {embedded_data};
const threshold = DATA.threshold;
document.querySelector('#threshold').textContent = threshold.toFixed(1);
const filters = {{ escalation: '', person: '', month: '', project: '', overOnly: false }};
const byId = id => document.getElementById(id);
const esc = value => String(value ?? '').replace(/[&<>"']/g, char => ({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[char]));
const format = value => Number(value).toFixed(2).replace(/\\.00$/, '');
const allWorkload = new Map(DATA.personMonths.map(row => [`${{row.email}}|${{row.month}}`, row.total_allocation]));
const allAssignments = new Map();
for (const row of DATA.contributions) {{
  const key = `${{row.email}}|${{row.month}}`;
  const item = allAssignments.get(key) || new Map();
  const projectKey = `${{row.project}}|${{row.escalation_path || 'Not specified'}}`;
  item.set(projectKey, (item.get(projectKey) || 0) + row.allocation);
  allAssignments.set(key, item);
}}
function option(select, value, label) {{ const node=document.createElement('option'); node.value=value; node.textContent=label; select.append(node); }}

const rowKeys = row => `${{row.sheet}}|${{row.row}}`;
const escalationByRow = new Map();
// Escalation data is supplied by the generator in the matching contribution rows.
for (const row of DATA.contributions) escalationByRow.set(rowKeys(row), row.escalation_path || 'Not specified');

function populateFilters() {{
  const escalationPaths = [...new Set(DATA.contributions.map(row => row.escalation_path || 'Not specified'))].sort();
  escalationPaths.forEach(value => option(byId('escalationFilter'), value, value));
  DATA.people.forEach(person => option(byId('personFilter'), person.email, `${{person.name}} (${{person.email}})`));
  DATA.months.forEach(month => option(byId('monthFilter'), month, month));
  DATA.projects.forEach(project => option(byId('projectFilter'), project, project));
}}

function matchesContribution(row) {{
  return (!filters.escalation || (row.escalation_path || 'Not specified') === filters.escalation)
    && (!filters.person || row.email === filters.person)
    && (!filters.month || row.month === filters.month)
    && (!filters.project || row.project === filters.project);
}}

function groupedPeople() {{
  const grouped = new Map();
  for (const row of DATA.contributions.filter(matchesContribution)) {{
    const key = `${{row.email}}|${{row.month}}`;
    if (!grouped.has(key)) grouped.set(key, {{ email:row.email, name:row.name || row.email, month:row.month, total:0, allTotal:allWorkload.get(key) || 0, projects:new Map(), escalation:new Set() }});
    const item = grouped.get(key); item.total += row.allocation;
    item.projects.set(row.project, (item.projects.get(row.project) || 0) + row.allocation);
    item.escalation.add(row.escalation_path || 'Not specified');
  }}
  return [...grouped.values()].filter(item => !filters.overOnly || item.allTotal > threshold + 1e-9)
    .sort((a,b) => b.allTotal-a.allTotal || a.month.localeCompare(b.month) || a.name.localeCompare(b.name));
}}

function renderCards(rows) {{
  const overloaded = rows.filter(item => item.allTotal > threshold + 1e-9);
  const excess = overloaded.reduce((sum,item) => sum + item.allTotal - threshold, 0);
  const people = new Set(rows.map(item => item.email));
  const max = rows.reduce((top,item) => Math.max(top,item.allTotal), 0);
  byId('cards').innerHTML = [
    ['Person-months shown', rows.length], ['People shown', people.size], ['Overloaded (all projects)', overloaded.length], ['Total excess FTE', format(excess)], ['Highest workload', format(max)]
  ].map(([label,value]) => `<article class="card"><div class="label">${{label}}</div><div class="value">${{value}}</div></article>`).join('');
}}

function renderPeople(rows) {{
  const body = rows.map(item => {{
    const projectList = [...item.projects.entries()].sort((a,b) => b[1]-a[1]).map(([project,value]) => `<span class="tag">${{esc(project)}} ${{format(value)}}</span>`).join('');
    const isOver = item.allTotal > threshold + 1e-9;
    return `<tr class="${{isOver?'over':''}}"><td>${{esc(item.month)}}</td><td><strong>${{esc(item.name)}}</strong><br><span class="muted">${{esc(item.email)}}</span></td><td>${{[...item.escalation].map(esc).join('<br>')}}</td><td class="total">${{format(item.total)}}</td><td class="total ${{isOver?'danger':''}}">${{format(item.allTotal)}}</td><td>${{projectList}}</td></tr>`;
  }}).join('');
  byId('people').innerHTML = `<h2>People workload</h2><p class="muted">Selected workload reflects the active filters. All-project workload remains the workload control used for overload warnings.</p>${{rows.length ? `<table><thead><tr><th>Month</th><th>Person</th><th>Escalation path</th><th>Selected workload</th><th>All-project workload</th><th>Project contributions</th></tr></thead><tbody>${{body}}</tbody></table>` : '<p class="empty">No matching workload records.</p>'}}`;
}}

function matrixPeople() {{
  const emails = new Set(DATA.contributions.filter(matchesContribution).map(row => row.email));
  return DATA.people.filter(person => emails.has(person.email)).sort((a,b) => a.name.localeCompare(b.name) || a.email.localeCompare(b.email));
}}

function matrixTooltip(person, month, total) {{
  const assignments = allAssignments.get(`${{person.email}}|${{month}}`) || new Map();
  const details = [...assignments.entries()].sort((a,b) => b[1]-a[1]).map(([key, allocation]) => {{
    const [project, escalation] = key.split('|');
    return `${{project}}: ${{format(allocation)}} (${{escalation}})`;
  }});
  return `${{person.name}} (${{person.email}})\n${{month}}\nAll-project workload: ${{format(total)}}\n\n${{details.join(String.fromCharCode(10))}}`;
}}

function bindMatrixTooltip(people) {{
  const popup = byId('matrixTooltip');
  const peopleByEmail = new Map(people.map(person => [person.email, person]));
  const show = (cell, event) => {{
    const person = peopleByEmail.get(cell.dataset.email);
    const month = cell.dataset.month;
    const total = allWorkload.get(`${{person.email}}|${{month}}`);
    const assignments = allAssignments.get(`${{person.email}}|${{month}}`) || new Map();
    const details = [...assignments.entries()].sort((a,b) => b[1]-a[1]).map(([key, allocation]) => {{
      const [project, escalation] = key.split('|');
      return `<div class="assignment">${{esc(project)}}: <strong>${{format(allocation)}}</strong><br><span>${{esc(escalation)}}</span></div>`;
    }}).join('<hr>');
    popup.innerHTML = `<div class="heading">${{esc(person.name)}} · ${{esc(month)}}<br>All-project workload: ${{format(total)}}</div>${{details}}`;
    popup.style.display = 'block';
    move(event);
  }};
  const move = event => {{
    const left = Math.min(event.clientX + 14, window.innerWidth - 430);
    const top = Math.min(event.clientY + 14, window.innerHeight - 220);
    popup.style.left = `${{Math.max(8, left)}}px`;
    popup.style.top = `${{Math.max(8, top)}}px`;
  }};
  document.querySelectorAll('#matrix td.load').forEach(cell => {{
    cell.addEventListener('mouseenter', event => show(cell, event));
    cell.addEventListener('mousemove', move);
    cell.addEventListener('mouseleave', () => {{ popup.style.display = 'none'; }});
  }});
}}

function renderMatrix() {{
  const people = matrixPeople();
  const months = filters.month ? [filters.month] : DATA.months;
  const head = months.map(month => `<th>${{esc(month)}}</th>`).join('');
  const body = people.map(person => {{
    const cells = months.map(month => {{
      const total = allWorkload.get(`${{person.email}}|${{month}}`);
      if (total === undefined) return '<td class="no-load">—</td>';
      const state = total > threshold + 1e-9 ? 'over' : total >= threshold * .8 ? 'watch' : '';
      return `<td class="load ${{state}}" data-email="${{esc(person.email)}}" data-month="${{esc(month)}}">${{format(total)}}</td>`;
    }}).join('');
    return `<tr><td class="person"><strong>${{esc(person.name)}}</strong><br><span class="muted">${{esc(person.email)}}</span></td>${{cells}}</tr>`;
  }}).join('');
  byId('matrix').innerHTML = `<h2>Workload matrix</h2><p class="muted">Filters select the people shown. Each cell is the employee's all-project workload; hover over a cell to see every project and escalation path behind that total.</p>${{people.length ? `<table class="matrix"><thead><tr><th>Person</th>${{head}}</tr></thead><tbody>${{body}}</tbody></table>` : '<p class="empty">No people match the active filters.</p>'}}`;
  bindMatrixTooltip(people);
}}

function renderProjects() {{
  let rows = DATA.contributions.filter(matchesContribution);
  if (filters.overOnly) {{
    const overKeys = new Set(groupedPeople().map(item => `${{item.email}}|${{item.month}}`));
    rows = rows.filter(item => overKeys.has(`${{item.email}}|${{item.month}}`));
  }}
  const grouped = new Map();
  for (const row of rows) {{
    const key = `${{row.project}}|${{row.month}}|${{row.email}}`;
    const item = grouped.get(key) || {{...row, total:0}}; item.total += row.allocation; grouped.set(key,item);
  }}
  const result = [...grouped.values()].sort((a,b) => a.project.localeCompare(b.project) || a.month.localeCompare(b.month) || b.total-a.total);
  const body = result.map(row => `<tr><td>${{esc(row.project)}}</td><td>${{esc(row.month)}}</td><td><strong>${{esc(row.name || row.email)}}</strong><br><span class="muted">${{esc(row.email)}}</span></td><td>${{esc(row.escalation_path || 'Not specified')}}</td><td>${{esc(row.function)}}</td><td class="total">${{format(row.total)}}</td></tr>`).join('');
  byId('projects').innerHTML = `<h2>Project assignments</h2>${{result.length ? `<table><thead><tr><th>Project</th><th>Month</th><th>Person</th><th>Escalation path</th><th>Function</th><th>Allocation</th></tr></thead><tbody>${{body}}</tbody></table>` : '<p class="empty">No matching project assignments.</p>'}}`;
}}

function renderMissing() {{
  const rows = DATA.missingEmails.filter(row => (!filters.project || row.project === filters.project) && (!filters.escalation || (row.escalation_path || 'Not specified') === filters.escalation));
  const body = rows.map(row => `<tr><td>${{esc(row.project)}}</td><td>${{esc(row.name || '—')}}</td><td>${{esc(row.function || '—')}}</td><td>${{esc(row.escalation_path || 'Not specified')}}</td><td>${{esc(row.email_cell || '—')}}</td><td>${{esc(row.monthly_allocations)}}</td></tr>`).join('');
  byId('missing').innerHTML = `<h2>TBH / rows without parseable email</h2>${{rows.length ? `<table><thead><tr><th>Project</th><th>Name</th><th>Function</th><th>Escalation path</th><th>Email cell</th><th>Monthly allocations</th></tr></thead><tbody>${{body}}</tbody></table>` : '<p class="empty">No matching rows.</p>'}}`;
}}

function render() {{ const rows=groupedPeople(); renderCards(rows); renderPeople(rows); renderMatrix(); renderProjects(); renderMissing(); }}
function wire(id, key) {{ byId(id).addEventListener('change', event => {{ filters[key]=event.target.type==='checkbox' ? event.target.checked : event.target.value; render(); }}); }}
populateFilters(); wire('escalationFilter','escalation'); wire('personFilter','person'); wire('monthFilter','month'); wire('projectFilter','project'); wire('overOnly','overOnly');
byId('reset').addEventListener('click', () => {{ for (const key of Object.keys(filters)) filters[key]=key==='overOnly'?false:''; byId('escalationFilter').value=''; byId('personFilter').value=''; byId('monthFilter').value=''; byId('projectFilter').value=''; byId('overOnly').checked=false; render(); }});
document.querySelectorAll('[data-tab]').forEach(button => button.addEventListener('click', () => {{ document.querySelectorAll('[data-tab]').forEach(node=>node.classList.toggle('active',node===button)); document.querySelectorAll('.tab').forEach(node=>node.classList.toggle('hidden',node.id!==button.dataset.tab)); }}));
render();
</script>
</body>
</html>"""


def enrich_with_escalation(workbook_path: Path, report_dir: Path) -> None:
    """Add the resource-row Escalation Path to report CSVs without editing source Excel."""
    import openpyxl

    from analyze_commitment import EXCLUDED_SHEETS

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    escalation_by_row = {}
    for sheet in workbook.worksheets:
        if sheet.title in EXCLUDED_SHEETS or sheet.cell(1, 1).value != "Project:":
            continue
        header_row = next((row for row in range(1, min(sheet.max_row, 20) + 1)
                           if any(str(sheet.cell(row, column).value or "").strip().lower() == "engineer email"
                                  for column in range(1, min(sheet.max_column, 20) + 1))), None)
        if not header_row:
            continue
        headers = {str(sheet.cell(header_row, column).value or "").strip().lower(): column
                   for column in range(1, sheet.max_column + 1)}
        escalation_column = headers.get("escalatin path (email)") or headers.get("escalation path (email)")
        if not escalation_column:
            continue
        for row in range(header_row + 1, sheet.max_row + 1):
            value = sheet.cell(row, escalation_column).value
            escalation_by_row[(sheet.title, str(row))] = str(value or "").strip() or "Not specified"

    for filename in ("allocation_contributions.csv", "rows_without_email.csv"):
        path = report_dir / filename
        rows = read_csv(path)
        if not rows:
            continue
        for row in rows:
            row["escalation_path"] = escalation_by_row.get((row["sheet"], row["row"]), "Not specified")
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=[*rows[0].keys()])
            writer.writeheader()
            writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Updated .xlsx or .xlsm commitment workbook")
    parser.add_argument("--report-dir", type=Path, default=Path("commitment_analysis"))
    parser.add_argument("--output", type=Path, default=Path("commitment_analysis/commitment_dashboard.html"))
    parser.add_argument("--threshold", type=float, default=1.0)
    args = parser.parse_args()

    analyze(args.workbook, args.report_dir, args.threshold)
    enrich_with_escalation(args.workbook, args.report_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(build_html(dashboard_data(args.report_dir, args.threshold), args.workbook.name), encoding="utf-8")
    print(f"Dashboard: {args.output}")


if __name__ == "__main__":
    main()
